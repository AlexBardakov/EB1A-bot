# app/core/ingest.py
import os
from datetime import datetime

# Библиотеки для текста
import docx
from pypdf import PdfReader

from sqlalchemy.orm import Session
from app.storage.models import Document, DocumentVersion, DocumentStatus, ChatState

# Папка для сохранения файлов (будет создана в корне проекта)
UPLOAD_DIR = os.path.join(os.getcwd(), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)


def extract_text_from_file(file_path: str, mime_type: str) -> str:
    """Извлекает чистый текст из PDF или DOCX."""
    text = ""
    try:
        if "pdf" in mime_type or file_path.endswith(".pdf"):
            reader = PdfReader(file_path)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"

        elif "word" in mime_type or "document" in mime_type or file_path.endswith(".docx"):
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"

        else:
            return "[Error: Unsupported file format for text extraction]"

    except Exception as e:
        return f"[Error extracting text: {e}]"

    return text.strip()


def save_document_upload(
        session: Session,
        chat_id: str,
        file_name: str,
        file_data: bytes,
        mime_type: str
) -> str:
    """
    Сохраняет файл, парсит текст, создает/обновляет запись в БД.
    """
    # 1. Проверка кейса
    cs = session.query(ChatState).filter(ChatState.chat_id == chat_id).one_or_none()
    if not cs or not cs.active_case_id:
        return "⚠️ Сначала выберите кейс через `/case use ...`"

    # 2. Сохранение на диск
    # Добавляем timestamp, чтобы файлы физически не перезаписывались
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    unique_name = f"{timestamp}_{file_name}"
    file_path = os.path.join(UPLOAD_DIR, unique_name)

    with open(file_path, "wb") as f:
        f.write(file_data)

    # 3. Экстракция текста
    extracted_text = extract_text_from_file(file_path, mime_type)
    if not extracted_text:
        extracted_text = "[No text extracted or empty file]"

    # 4. Работа с БД
    # Ищем документ по имени в текущем кейсе
    existing_doc = (
        session.query(Document)
        .filter(Document.case_id == cs.active_case_id, Document.title == file_name)
        .one_or_none()
    )

    if existing_doc:
        doc = existing_doc
        # Считаем текущие версии, чтобы показать красивый номер
        ver_num = len(doc.versions) + 1
        msg_prefix = f"🔄 Документ '{file_name}' найден. Добавляю новую версию (v{ver_num})..."
    else:
        doc = Document(
            case_id=cs.active_case_id,
            title=file_name,
            doc_type="unknown",
            status=DocumentStatus.draft
        )
        session.add(doc)
        session.flush()  # Получаем ID
        msg_prefix = f"✅ Создан новый документ '{file_name}'."

    # Создаем версию
    version = DocumentVersion(
        document_id=doc.id,
        storage_url=file_path,
        text_extract=extracted_text,
        created_by=f"user_{chat_id}",
        notes=f"Uploaded via Telegram at {timestamp}"
    )
    session.add(version)
    session.flush()

    # Обновляем ссылку на текущую версию
    doc.current_version_id = version.id
    doc.updated_at = datetime.utcnow()

    return (
        f"{msg_prefix}\n"
        f"Символов распознано: {len(extracted_text)}\n"
        f"Теперь можно проверить: `/review {file_name}`"
    )